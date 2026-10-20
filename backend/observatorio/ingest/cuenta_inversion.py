"""Conector Cuenta de Inversión (Contaduría General de la Nación).

Cuadro 35 del Tomo I: Cuenta Ahorro-Inversión-Financiamiento del Sector Público Nacional no
Financiero, por nivel institucional, con eliminaciones de consolidación. Verificado 2026-10-19:
  página  https://www.argentina.gob.ar/economia/sechacienda/cgn/cuenta/2025/cuenta-de-inversion-2025-tomo-i/gestion-financiera-consolidada
  archivo https://www.argentina.gob.ar/sites/default/files/cuadro35-tomoi-2025.xlsx
Importes en pesos, guardados en el xlsx como números de doble precisión (≈16 cifras
significativas): para montos de 10^14 $ el último centavo no es exacto en el propio archivo.
Se conserva el repr del valor y se convierte a centavos sin redondear más allá de eso.

Criterio de registro según la CGN (gestion-financiera-consolidada_aspectos_generales_y_metodologicos-tomoi-2025.pdf):
devengado para erogaciones y recaudado/percibido para recursos; las transferencias entre entes
del universo se toman devengadas de ambos lados para poder eliminarlas.
"""
from __future__ import annotations

import re
import sqlite3
import unicodedata
from decimal import ROUND_HALF_EVEN, Decimal
from pathlib import Path

from ..db import now_iso
from .base import DatasetDef, IngestError, SourceDef

CGN = SourceDef(
    source_id="cgn_cuenta_inversion",
    organismo="Contaduría General de la Nación (Secretaría de Hacienda, MECON)",
    base_url="https://www.argentina.gob.ar/economia/sechacienda/cgn/cuentainversion",
    licencia=None,
    frecuencia_publicada="anual (cierre del ejercicio, junio del año siguiente)",
    notas="Cuenta de Inversión remitida al Congreso (Ley 24.156, art. 95).",
)
BASE_FILES = "https://www.argentina.gob.ar/sites/default/files/"
UNIVERSO = "Sector Público Nacional no Financiero"
BASE_CONTABLE = "devengado (gastos) / percibido (recursos); transferencias intra-SPNF devengadas"

# Columna -> (subsector, rótulo esperado en filas 6-8 del encabezado, normalizado)
COLUMNAS = {
    "B": ("adm_central", "administracion central"),
    "C": ("org_descentralizados", "organismos descentralizados"),
    "D": ("inst_seguridad_social", "instituciones de seguridad social"),
    "E": ("total_an", "total"),
    "F": ("universidades", "universidades nacionales"),
    "G": ("empresas", "empresas publicas"),
    "H": ("fondos_fiduciarios", "fondos fiduciarios"),
    "I": ("entes_publicos", "entes publicos"),
    "J": ("eliminaciones", "eliminacion"),
    "K": ("total_spnf", "total sector publico no financiero"),
}
SUBSECTORES_SUMA = ("total_an", "universidades", "empresas", "fondos_fiduciarios", "entes_publicos")
AN_PARTES = ("adm_central", "org_descentralizados", "inst_seguridad_social")
# Tolerancia de las identidades: el archivo guarda dobles; con 10^14 $ el error de representación
# es ~0,02 $ por celda (máximo observado en 2025: 0,021 $ en total AN = suma de sus partes).
TOLERANCIA_CENTAVOS = 10
ROMANOS = ("I", "II", "III", "IV", "V", "VI", "VII", "VIII", "IX", "X", "XI", "XII", "XIII", "XIV")


def cuadro35_def(ejercicio: int) -> DatasetDef:
    return DatasetDef(
        dataset_id=f"cgn_ci_{ejercicio}_cuadro35", source=CGN,
        nombre=f"Cuenta de Inversión {ejercicio} · Tomo I · Cuadro 35: Cuenta AIF del SPNF por nivel institucional",
        url=f"{BASE_FILES}cuadro35-tomoi-{ejercicio}.xlsx", universo=UNIVERSO,
        nivel_gobierno="nacional", unidad="pesos corrientes", cobertura_temporal=str(ejercicio),
        granularidad="anual, por subsector institucional",
        metodologia=BASE_CONTABLE + ". Incluye columna de eliminaciones de consolidación.")


def _norm(s: str) -> str:
    s = unicodedata.normalize("NFKD", s).encode("ascii", "ignore").decode().lower()
    return " ".join(re.sub(r"[^a-z0-9 ]", " ", s).split())


def _slug(s: str) -> str:
    return _norm(s).replace(" ", "_")


def to_centavos(v) -> tuple[int, str]:
    d = Decimal(repr(v)) if isinstance(v, float) else Decimal(v)
    return int((d * 100).quantize(Decimal(1), rounding=ROUND_HALF_EVEN)), repr(v)


def _nivel(label: str) -> tuple[int, str]:
    """Jerarquía según la sangría del cuadro: 'I ...' (0), '- x' (1), '. x' (2), '.. x' (3)."""
    t = label.strip()
    m = re.match(r"^(X{0,3}(?:IX|IV|V?I{0,3}))\s+(.*)$", t)
    if m and m.group(1) in ROMANOS:
        return 0, m.group(1)
    if t.startswith(".."):
        return 3, t[2:].strip()
    if t.startswith("."):
        return 2, t[1:].strip()
    if t.startswith("-"):
        return 1, t[1:].strip()
    raise IngestError(f"Cuadro 35: rótulo sin jerarquía reconocible: {label!r}")


def parse_cuadro35(path: Path, ejercicio: int):
    import openpyxl
    ws = openpyxl.load_workbook(path, data_only=True).active
    head = {c: _norm(str(ws[c].value or "")) for c in ("K1", "A2", "A3", "A4", "K5")}
    expected = {"K1": "cuadro nro 35", "A2": "sector publico nacional no financiero",
                "A3": f"ejercicio {ejercicio}", "A4": "cuenta ahorro inversion financiamiento",
                "K5": "en pesos"}
    for c, exp in expected.items():
        if head[c] != exp:
            raise IngestError(f"Cuadro 35: encabezado {c}={head[c]!r}, se esperaba {exp!r}")
    for col, (_, rot) in COLUMNAS.items():
        got = _norm(" ".join(str(ws[f"{col}{r}"].value or "") for r in (6, 7, 8)))
        if col in "BCD":  # 'Administración Nacional' está en B6, encima de las tres partes
            got = got.replace("administracion nacional ", "", 1) if col == "B" else got
        if got != rot:
            raise IngestError(f"Cuadro 35: columna {col} rotulada {got!r}, se esperaba {rot!r}")

    cells, path_stack, seen = [], [], set()
    for r in range(9, ws.max_row + 1):
        label = ws[f"A{r}"].value
        if not label or not str(label).strip():
            continue  # filas separadoras (algunas traen un 0 suelto en K: no es dato)
        nivel, texto = _nivel(str(label))
        path_stack = path_stack[:nivel] + [texto if nivel == 0 else _slug(texto)]
        cid = ".".join(path_stack)
        if cid in seen:
            raise IngestError(f"Cuadro 35: concepto duplicado {cid}")
        seen.add(cid)
        for col, (sub, _) in COLUMNAS.items():
            v = ws[f"{col}{r}"].value
            if v is None:
                estado, cent, orig = "vacia", None, None
            elif isinstance(v, str) and v.startswith("#"):
                estado, cent, orig = "error_formula", None, v
            elif isinstance(v, (int, float)):
                estado, (cent, orig) = "ok", to_centavos(v)
            else:
                raise IngestError(f"Cuadro 35: valor no numérico en {col}{r}: {v!r}")
            cells.append(dict(celda=f"{col}{r}", concepto_id=cid, concepto=str(label).strip(),
                              subsector=sub, importe_centavos=cent, importe_original=orig, estado=estado))
    if not any(c["concepto_id"] == "XI" for c in cells):
        raise IngestError("Cuadro 35: falta la fila XI RESULTADO FINANCIERO")
    return cells


def identity_checks(cells: list[dict]) -> list[dict]:
    """Identidades contables del cuadro. Celda vacía = no participa (no se toma como 0)."""
    val = {(c["concepto_id"], c["subsector"]): c["importe_centavos"] for c in cells}
    conceptos = list(dict.fromkeys(c["concepto_id"] for c in cells))
    checks = []

    def add(nombre, calc, pub):
        if calc is None or pub is None:
            checks.append(dict(nombre=nombre, calculado=calc, publicado=pub, diferencia=None, conclusion="sin_dato"))
        else:
            d = calc - pub
            checks.append(dict(nombre=nombre, calculado=calc, publicado=pub, diferencia=d,
                               conclusion="coincide" if abs(d) <= TOLERANCIA_CENTAVOS else "discrepa"))

    def s(*keys, signs=None):
        signs = signs or [1] * len(keys)
        vals = [val.get(k) for k in keys]
        if any(v is None for v in vals):
            return None
        return sum(sg * v for sg, v in zip(signs, vals))

    for cid in conceptos:
        parts = [val.get((cid, p)) for p in AN_PARTES]
        if all(p is not None for p in parts) and val.get((cid, "total_an")) is not None:
            add(f"{cid}: total AN = central + descentralizados + seg. social", sum(parts), val[(cid, "total_an")])
        tot = val.get((cid, "total_spnf"))
        present = [val.get((cid, x)) for x in SUBSECTORES_SUMA if val.get((cid, x)) is not None]
        if tot is not None and present:
            calc = sum(present) - (val.get((cid, "eliminaciones")) or 0)
            add(f"{cid}: total SPNF = Σ subsectores − eliminaciones", calc, tot)
    for sub in [v[0] for v in COLUMNAS.values()]:
        add(f"III [{sub}] = I − II", s(("I", sub), ("II", sub), signs=[1, -1]), val.get(("III", sub)))
        add(f"VI [{sub}] = I + IV", s(("I", sub), ("IV", sub)), val.get(("VI", sub)))
        add(f"VII [{sub}] = II + V", s(("II", sub), ("V", sub)), val.get(("VII", sub)))
        add(f"VIII [{sub}] = VI − VII", s(("VI", sub), ("VII", sub), signs=[1, -1]), val.get(("VIII", sub)))
        if sub != "eliminaciones":
            ix, x = val.get(("IX", sub)), val.get(("X", sub))
            calc = None if val.get(("VIII", sub)) is None else val[("VIII", sub)] + (ix or 0) - (x or 0)
            add(f"XI [{sub}] = VIII + IX − X", calc, val.get(("XI", sub)))
            add(f"XIV [{sub}] = XII − XIII", s(("XII", sub), ("XIII", sub), signs=[1, -1]),
                val.get(("XIV", sub)))
            xi = val.get(("XI", sub))
            add(f"XIV [{sub}] = −XI (el financiamiento neto cubre el resultado)",
                None if xi is None else -xi, val.get(("XIV", sub)))
    # Consolidación: las contribuciones figurativas se eliminan contra los gastos figurativos.
    add("Eliminaciones: contribuciones figurativas (IX) = gastos figurativos (X)",
        val.get(("IX", "eliminaciones")), val.get(("X", "eliminaciones")))
    add("Eliminaciones: transferencias corrientes recibidas = otorgadas",
        val.get(("I.transferencias_corrientes", "eliminaciones")),
        val.get(("II.transferencias_corrientes", "eliminaciones")))
    add("Eliminaciones: no alteran el resultado económico (III)", 0, val.get(("III", "eliminaciones")))
    return checks


def make_loader(ejercicio: int, ds: DatasetDef):
    def load(con: sqlite3.Connection, path: Path, run_id: int):
        cells = parse_cuadro35(path, ejercicio)
        con.execute("DELETE FROM aif_observation WHERE dataset_id=?", (ds.dataset_id,))
        con.executemany(
            "INSERT INTO aif_observation VALUES (?,?,?,?,?,?,?,?,?,?,?,?,?)",
            [(ds.dataset_id, run_id, ejercicio, "35", c["celda"], c["concepto_id"], c["concepto"],
              c["subsector"], UNIVERSO, BASE_CONTABLE, c["importe_centavos"], c["importe_original"],
              c["estado"]) for c in cells])
        prefix = f"ci35:{ejercicio}:"
        con.execute("DELETE FROM reconciliation WHERE comparacion LIKE ?", (prefix + "%",))
        for ch in identity_checks(cells):
            con.execute("INSERT INTO reconciliation (run_id, comparacion, total_calculado_centavos, "
                        "total_publicado_centavos, diferencia_centavos, tolerancia_centavos, conclusion, "
                        "checked_at) VALUES (?,?,?,?,?,?,?,?)",
                        (run_id, prefix + ch["nombre"], ch["calculado"], ch["publicado"], ch["diferencia"],
                         TOLERANCIA_CENTAVOS, ch["conclusion"], now_iso()))
        errores = sum(c["estado"] == "error_formula" for c in cells)
        return len(cells), (f"Cuenta de Inversión {ejercicio} (cierre)"
                            + (f"; {errores} celdas con error de fórmula en el archivo oficial" if errores else ""))
    return load


def metodologia_def(ejercicio: int) -> DatasetDef:
    """PDF de criterios de la Gestión Financiera Consolidada (fuente de citas metodológicas)."""
    return DatasetDef(
        dataset_id=f"cgn_ci_{ejercicio}_metodologia_spnf", source=CGN,
        nombre=f"Cuenta de Inversión {ejercicio} · Gestión Financiera Consolidada: aspectos generales y metodológicos",
        url=f"{BASE_FILES}gestion-financiera-consolidada_aspectos_generales_y_metodologicos-tomoi-{ejercicio}.pdf",
        universo=UNIVERSO, nivel_gobierno="nacional", unidad="documento (texto)",
        cobertura_temporal=str(ejercicio), granularidad="documento",
        metodologia="Documento metodológico; no contiene series. Se usa para verificar citas textuales.")


def load_documento(con: sqlite3.Connection, path: Path, run_id: int):
    from .boletin_fiscal import page_texts
    pages = page_texts(path)
    if "GESTIÓN FINANCIERA CONSOLIDADA" not in " ".join(pages[0].split()).upper():
        raise IngestError("El PDF metodológico de la CGN no tiene el título esperado")
    return len(pages), "Cuenta de Inversión (documento metodológico)"
